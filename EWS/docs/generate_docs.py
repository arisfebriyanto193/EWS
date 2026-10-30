import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = "/home/aris/Dokumen/projeck/EWS_new/EWS/EWS/docs"
IMG_DIR = os.path.join(DOCS_DIR, "images")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 81, 123) # Slate Blue
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(53, 95, 142)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
    return p

def add_paragraph_styled(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10)
        r_pre.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.italic = italic
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, title, text, bg_color="F0FDF4", border_color="22C55E"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border styling left only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + "\n")
    r1.font.name = 'Arial'
    r1.font.size = Pt(10.5)
    r1.bold = True
    r1.font.color.rgb = RGBColor(21, 128, 61) if border_color == "22C55E" else RGBColor(185, 28, 28)
    
    r2 = p.add_run(text)
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy Blue
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    # Column widths
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_image_with_caption(doc, img_rel_path, caption, width_inch=5.8):
    full_path = os.path.join(IMG_DIR, img_rel_path)
    if os.path.exists(full_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(full_path, width=Inches(width_inch))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(f"Gambar: {caption}")
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(8.5)
        run_cap.italic = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139)

def generate_technical_doc():
    doc = docx.Document()
    
    # Page setup (A4, 0.75 in margins)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Dokumen Teknis Rekayasa Sistem EWS & Perangkap Hama Terintegrasi")
        f_run.font.name = 'Arial'
        f_run.font.size = Pt(8)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("DOKUMEN SPESIFIKASI TEKNIS & ARSITEKTUR REKAYASA")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(17)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Sistem Landslide Early Warning System (EWS) 4 Titik & Perangkap Hama Surya 5 Unit\nIntegrasi Hardware PCB, Firmware ESP32-S3 FreeRTOS, Backend Express WS, & Frontend Next.js")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Table
    meta_headers = ["Parameter Sistem", "Spesifikasi Implementasi & Detail Teknis"]
    meta_data = [
        ["Target Pengguna", "Teknisi Elektronika, Field Engineer, Developer IoT, BPBD System Admin"],
        ["Mikrokontroler Node", "ESP32-S3 (Dual-Core Xtensa LX7 @ 240 MHz, 8MB PSRAM, 16MB Flash)"],
        ["Sensor Pemantau", "Inclinometer MPU-6050 Dual-Axis, Tipping Bucket 0.2mm, Getaran 801S, INA226, Kelembapan Tanah"],
        ["Sistem Catu Daya", "Panel Surya 30 Wp Monocrystalline + Baterai VRLA 12V + Smart Charge Controller"],
        ["Sistem Alarm Lapangan", "Sirine Otomotif 12V (115-120 dB), Lampu Strobo Darurat Merah/Kuning, Driver Relay Terisolasi"],
        ["Konektivitas Transmisi", "Modem Seluler GSM/4G SIMCom A7670C / Wi-Fi Dual-Link (Fallback Lokal)"],
        ["Protokol Komunikasi", "WebSocket Duplex (Port 5000 / ws:// / wss://) & REST API HTTP JSON"],
        ["Backend Gateway", "Node.js (Express), ws Engine, Evaluasi Ambang Batas Real-Time, Bot Telegram"],
        ["Database Relasional", "MySQL / MariaDB (Tabel Node, Histori Log Sensor, Perangkap Hama, Audit Trail)"],
        ["Frontend Dashboard", "Next.js 16 (App Router), React 19, TypeScript, Tailwind CSS, Leaflet Satellite GIS"],
    ]
    add_table_data(doc, meta_headers, meta_data, col_widths=[2.2, 4.5])

    # Section 1
    add_header_styled(doc, "1. Arsitektur Global & Alur Data Transmisi", level=1)
    add_paragraph_styled(doc, "Sistem dirancang dengan arsitektur telemetri terdistribusi yang memadukan keandalan respons tepi (Edge Computing) dan kemudahan pemantauan awan (Cloud Dashboard). Node EWS membaca kondisi lereng secara periodik dan mengevaluasi status bahaya secara mandiri tanpa ketergantungan koneksi internet.")
    
    add_paragraph_styled(doc, "Alur data mengalir dari sensor analog dan digital ke mikrokontroler ESP32-S3. Firmware memfilter noise bacaan dengan moving-average, kemudian membungkus data dalam format JSON terkompresi. Data dikirimkan secara langsung melalui protokol WebSocket dua arah ke Gateway Server Express.js, dicatat ke basis data MySQL, dan disiarkan secara instan ke seluruh antarmuka web monitoring.")
    
    add_callout(doc, "PENGAMANAN TINGKAT TINGGI: FAIL-SAFE LOKAL", 
                "Jika jaringan seluler GSM atau server backend terputus, algoritma di dalam mikrokontroler ESP32-S3 tetap menjalankan proteksi keselamatan. Sirine 115 dB dan lampu strobo akan langsung aktif di tiang begitu kemiringan tanah atau intensitas hujan melewati ambang batas bahaya.",
                bg_color="FEF2F2", border_color="EF4444")

    # Section 2
    add_header_styled(doc, "2. Desain Elektronika & Rangkaian Skematik PCB", level=1)
    add_paragraph_styled(doc, "Rangkaian PCB EWS-45 dirancang dengan mengutamakan isolasi interferensi frekuensi tinggi dari modem GSM dan perlindungan lonjakan arus aktuator sirine 12V.")
    
    add_image_with_caption(doc, "schematic_ews45-1.png", "Skematik Rangkaian Elektronika PCB EWS-45 (Sensor Bus, MCU, Power Rail & Relay Driver)", width_inch=6.5)

    add_header_styled(doc, "2.1 Konfigurasi Pinout Mikrokontroler ESP32-S3", level=2)
    pin_headers = ["Simbol Net", "GPIO ESP32-S3", "Tipe I/O", "Komponen Terkoneksi", "Fungsi & Operasional"]
    pin_data = [
        ["IN2", "GPIO 46", "Digital Out", "Driver Relay Sirine & Strobo", "Active-HIGH. Memicu transistor driver relay 12V sirine"],
        ["Dhujan", "GPIO 10", "Digital In", "Sensor Hujan Tipping Bucket", "Interrupt digital; 1 pulsa ayunan = 0.20 mm curah hujan"],
        ["Dgetar", "GPIO 7", "Digital In", "Sensor Getaran Shock 801S", "Interrupt input; menghitung frekuensi getaran mikro lereng"],
        ["Ktanah", "GPIO 17", "Analog In", "Sensor Kelembapan Tanah", "ADC2_CH6; membaca resistansi/kapasitansi kelembapan"],
        ["SDAoled", "GPIO 8", "I2C Data", "Layar OLED 1.3\" (0x3C)", "Menampilkan status IP, tegangan, dan sensor lokal"],
        ["SCLoled", "GPIO 9", "I2C Clock", "Layar OLED 1.3\" (0x3C)", "Bus I2C Hardware 0 @ 400 kHz"],
        ["SDAMpu", "GPIO 11", "I2C Data", "Inclinometer MPU-6050", "Bus I2C Terisolasi untuk sudut kemiringan lereng X/Y"],
        ["SCLMpu", "GPIO 12", "I2C Clock", "Inclinometer MPU-6050", "Menghindari hang bus jika sensor lain mengalami gangguan"],
        ["SCL2/SDA", "GPIO 13", "I2C Data", "INA226 Power Monitor", "Sensor tegangan bus baterai dan konsumsi arus beban"],
        ["SDA2/SCL", "GPIO 14", "I2C Clock", "INA226 Power Monitor", "Bus I2C Hardware 1"],
        ["SDAds", "GPIO 15", "I2C Data", "RTC DS3231 Presisi Tinggi", "Waktu akurat untuk penanggalan offline (*blackbox*)"],
        ["SCLds", "GPIO 16", "I2C Clock", "RTC DS3231 Presisi Tinggi", "Dilengkapi baterai cadangan CR2032"],
        ["CS", "GPIO 6", "SPI CS", "Modul MicroSD Card", "Penyimpanan histori log berkala ke kartu memori"],
        ["SCK / MOSI / MISO", "GPIO 5/4/2", "SPI Bus", "Modul MicroSD Card", "Komunikasi SPI berkecepatan tinggi"],
    ]
    add_table_data(doc, pin_headers, pin_data, col_widths=[1.0, 1.1, 0.9, 1.8, 2.0])

    # Section 3
    add_header_styled(doc, "3. Logika Firmware ESP32-S3 (FreeRTOS Multi-Tasking)", level=1)
    add_paragraph_styled(doc, "Firmware memanfaatkan sistem operasi waktu nyata FreeRTOS dengan pemisahan tugas inti dual-core:")
    add_paragraph_styled(doc, "Tugas jaringan memelihara soket koneksi WebSocket client, memproses handshake TLS/SSL, mengirimkan paket telemetri dengan pacing buffer, serta mendengarkan pesan perintah kontrol.", bold_prefix="• Core 0 (Network Task): ")
    add_paragraph_styled(doc, "Tugas akuisisi sensor melakukan sampling berkala I2C, kalkulasi kemiringan sudut MPU-6050, perhitungan moving-average kelembapan tanah, penghitungan pulsa rain gauge, dan evaluasi logika status bahaya.", bold_prefix="• Core 1 (Sensor & Safety Task): ")
    add_paragraph_styled(doc, "Sinkronisasi antar-tugas dikelola menggunakan mutex `wsMutex` untuk soket jaringan, `dataMutex` untuk snapshot sensor bersama, dan critical section `portENTER_CRITICAL` untuk interupsi sensor getaran 801S.", bold_prefix="• Sinkronisasi Thread-Safe: ")

    add_header_styled(doc, "3.1 Tabel Kode Status & Anomali Sistem", level=2)
    code_headers = ["Kode", "Status Logika", "Kondisi Ambang Batas", "Respon Sirine & Strobo", "Notifikasi Telegram"]
    code_data = [
        ["OK", "Normal / Aman", "Semua sensor dalam batas normal", "Mati (Standby)", "Tidak ada"],
        ["A1", "Baterai Lemah", "Tegangan aki < 11.20 V", "Mati (Standby)", "Pesan peringatan teknisi"],
        ["A2", "Bahaya Longsor", "Sudut Kemiringan lereng >= Ambang Bahaya", "MENYALA PENUH (115 dB)", "Peringatan Darurat EVAKUASI"],
        ["A3", "Siaga Longsor", "Sudut Kemiringan lereng >= Ambang Siaga", "Mati / Strobo berkedip", "Peringatan Waspada Level 2"],
        ["A4", "Bahaya Hujan", "Curah hujan >= 50.0 mm/jam (Ekstrem)", "MENYALA PENUH (115 dB)", "Peringatan Banjir & Longsor"],
        ["A5", "Siaga Hujan", "Curah hujan >= 20.0 mm/jam (Lebat)", "Mati (Standby)", "Peringatan Siaga Hujan"],
        ["A6", "Bahaya Getaran", "Frekuensi getaran >= Ambang Bahaya", "MENYALA PENUH (115 dB)", "Peringatan Gempa/Retakan"],
        ["A7", "Siaga Tanah", "Kelembapan tanah >= 85% (Jenuh)", "Mati (Standby)", "Himbauan Patroli Lereng"],
        ["A8", "Filter Getaran", "Getaran sesaat terdeteksi", "Mati (Tahap Verifikasi)", "Tidak ada"],
    ]
    add_table_data(doc, code_headers, code_data, col_widths=[0.6, 1.2, 1.8, 1.7, 1.6])

    # Section 4
    add_header_styled(doc, "4. Arsitektur Backend (Node.js & WebSocket Gateway)", level=1)
    add_paragraph_styled(doc, "Backend bertindak sebagai sentral penerima telemetri dengan rute API REST dan port WebSocket dua arah:")
    add_paragraph_styled(doc, "Menyediakan endpoint otentikasi JWT (`/ews/api/auth/login`), manajemen titik EWS (`/ews/api/nodes`), pengaturan ambang batas (`/ews/api/nodes/:id/thresholds`), data histori (`/ews/api/logs`), dan audit trail (`/ews/api/alarms`).", bold_prefix="• Express REST API: ")
    add_paragraph_styled(doc, "Menerima koneksi persistent dari ESP32 dan Frontend browser. Menyediakan mekanisme verifikasi token, heartbeat ping/pong setiap 10 detik, dan broadcast otomatis setiap kali ada anomali sensor.", bold_prefix="• WebSocket Gateway Server: ")
    add_paragraph_styled(doc, "Jika terdeteksi status A2, A4, atau A6, server otomatis mengirimkan payload darurat ke bot Telegram yang langsung meneruskan broadcast ke grup siaga bencana desa.", bold_prefix="• Telegram Bot Dispatcher: ")

    # Section 5
    add_header_styled(doc, "5. Arsitektur Frontend Next.js 16 Web Monitoring", level=1)
    add_paragraph_styled(doc, "Frontend dibangun menggunakan Next.js 16 App Router dengan performa rendering tinggi dan antarmuka responsif.")
    
    add_image_with_caption(doc, "screenshot_dashboard.png", "Dashboard Utama Web Monitoring EWS & Perangkap Hama", width_inch=6.5)
    add_image_with_caption(doc, "screenshot_overview.png", "Pemetaan Peta Satelit Interaktif 4 Titik EWS dan 5 Perangkap Hama", width_inch=6.5)
    add_image_with_caption(doc, "screenshot_ews_detail.png", "Halaman Detail Sensor EWS: Inclinometer X/Y, Hujan, Baterai, dan Aktuator", width_inch=6.5)
    add_image_with_caption(doc, "screenshot_settings.png", "Panel Konfigurasi Ambang Batas Sensor, Kuota Telkomsel, & Akun", width_inch=6.5)

    # Section 6
    add_header_styled(doc, "6. Pemeliharaan Rutin & Kalibrasi Sensor", level=1)
    add_paragraph_styled(doc, "Lakukan kalibrasi tare zeroing pada sensor MPU-6050 saat pemasangan awal pada tiang dengan menekan tombol Kalibrasi Nol di Web Dashboard saat tiang terpasang tegak.", bold_prefix="• Kalibrasi Kemiringan: ")
    add_paragraph_styled(doc, "Bersihkan corong penampung air hujan dari daun kering atau debu setiap 30 hari sekali, serta uji ayunan mekanik tipping bucket menggunakan air 20 ml.", bold_prefix="• Pemeliharaan Rain Gauge: ")
    add_paragraph_styled(doc, "Pastikan kaca panel surya bersih dan pantau tegangan baterai pada dashboard (rentang sehat: 12.0V - 14.2V).", bold_prefix="• Pemeriksaan Catu Daya: ")

    output_path = os.path.join(DOCS_DIR, "Dokumen_Teknis_EWS_dan_Perangkap_Hama.docx")
    doc.save(output_path)
    print("Saved Technical Doc to:", output_path)

def generate_user_guide_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Buku Panduan Penggunaan Sistem EWS & Perangkap Hama - Untuk Masyarakat & Operator")
        f_run.font.name = 'Arial'
        f_run.font.size = Pt(8)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("BUKU PANDUAN PENGGUNAAN SISTEM")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(18)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(22, 101, 52) # Forest Green
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Peringatan Dini Bencana Tanah Longsor (EWS) & Perangkap Hama Pertanian Tenaga Surya\nPanduan Akses Web Monitoring, Arti Tanda Bahaya Sirine, & Protokol Kesiapsiagaan Warga")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Info Table
    info_headers = ["Bagian Panduan", "Keterangan Petunjuk"]
    info_data = [
        ["Tujuan Buku", "Memberikan pemahaman praktis kepada warga dan relawan mengenai cara kerja sistem EWS dan langkah keselamatan saat bahaya longsor."],
        ["Sasaran Pembaca", "Warga Masyarakat Sekitar Lereng, Perangkat Desa, Relawan Bencana, Petani, dan Operator Lapangan."],
        ["Perangkat Terpasang", "4 Titik Tiang EWS di Lereng Rawan & 5 Titik Perangkap Hama Surya di Areal Sawah/Kebun."],
        ["Akses Website", "Dapat dibuka gratis lewat Browser HP Android / iPhone atau Laptop."],
        ["Saluran Notifikasi", "Grup Telegram Resmi Siaga Bencana Desa & Suara Sirine Langsung di Tiang."],
    ]
    add_table_data(doc, info_headers, info_data, col_widths=[2.2, 4.5])

    # Section 1
    add_header_styled(doc, "1. Mengenal Alat di Lingkungan Desa Anda", level=1)
    add_paragraph_styled(doc, "Di wilayah desa kita, telah dipasang alat pelindung warga dan pertanian yang bekerja siang dan malam secara otomatis:")
    add_paragraph_styled(doc, "Tiang besi kokoh yang dipasang di tebing lereng. Tiang ini memiliki sensor pendeteksi kemiringan tanah, sensor curah hujan, sensor getaran, lampu strobo darurat, dan sirine bersuara sangat nyaring.", bold_prefix="1. Tiang EWS Longsor: ")
    add_paragraph_styled(doc, "Perangkat bertenaga surya yang dipasang di lahan sawah/kebun untuk membasmi serangga hama perusak tanaman menggunakan lampu pemikat otomatis tanpa racun kimia.", bold_prefix="2. Perangkap Hama Surya: ")

    # Section 2
    add_header_styled(doc, "2. Memahami Arti Bunyi Sirine & Lampu di Tiang EWS", level=1)
    add_paragraph_styled(doc, "Warga diharapkan mengenali tiga tingkatan status peringatan yang terlihat dan terdengar dari tiang EWS:")

    siren_headers = ["Tanda di Tiang EWS", "Tingkat Status", "Arti Bagi Warga", "Tindakan yang Harus Dilakukan"]
    siren_data = [
        ["Lampu Hijau Berkedip Tenang\n(Sirine Mati)", "AMAN", "Kondisi lereng sangat stabil, tidak ada gerakan tanah, curah hujan normal.", "Warga dapat bekerja dan beristirahat seperti biasa dengan tenang."],
        ["Lampu Kuning Berkedip Cepat\n(Sirine Mati)", "WASPADA / SIAGA", "Hujan lebat berkepanjangan, tanah lereng mulai jenuh air, ada getaran kecil.", "Tingkatkan kewaspadaan, siapkan tas dokumen penting, pantau lereng jika ada rembesan air keruh."],
        ["SIRINE BERBUNYI NYARING (115 dB)\nLampu Strobo Merah Berkedip", "BAHAYA / EVAKUASI", "TEBING TANAH BERGERAK ATAU RETAK! Ancaman longsor sedang terjadi!", "SEGERA KELUAR RUMAH! Berjalan cepat menjauhi lereng menuju Titik Kumpul Aman!"],
    ]
    add_table_data(doc, siren_headers, siren_data, col_widths=[1.6, 1.1, 2.0, 2.1])

    add_callout(doc, "CATATAN PENTING KESELAMATAN WARGA",
                "Sirine pada tiang EWS tetap berbunyi keras secara otomatis meskipun aliran listrik PLN padam atau sinyal HP hilang. Jika sirine berbunyi nyaring, jangan menunggu konfirmasi tambahan, segera ajak keluarga mengungsi ke tempat terbuka!",
                bg_color="FEF2F2", border_color="EF4444")

    # Section 3
    add_header_styled(doc, "3. Panduan Menggunakan Website Web Monitoring", level=1)
    add_paragraph_styled(doc, "Warga dan pengurus desa dapat memantau status keamanan desa secara langsung dari layar HP:")
    add_paragraph_styled(doc, "Buka aplikasi Google Chrome atau peramban lainnya di HP, lalu ketik alamat website monitoring desa.", bold_prefix="1. Membuka Website: ")
    add_paragraph_styled(doc, "Perhatikan 4 kotak informasi utama. Jika kotak bertuliskan AMAN (berwarna hijau), artinya wilayah aman. Jika berwarna oranye atau merah, perhatikan titik lereng yang disebutkan.", bold_prefix="2. Membaca Status: ")

    add_image_with_caption(doc, "screenshot_dashboard.png", "Tampilan Beranda Website: Memantau Kondisi 4 Titik EWS dan Cuaca", width_inch=6.5)
    
    add_header_styled(doc, "3.1 Melihat Peta Wilayah & Titik Rawan", level=2)
    add_paragraph_styled(doc, "Tekan menu Peta Satelit untuk melihat peta foto udara desa. Di peta ini, warga dapat melihat titik mana saja yang rawan dan lokasi jalur evakuasi aman terdekat.")
    add_image_with_caption(doc, "screenshot_overview.png", "Peta Satelit Interaktif Desa dengan Penanda Lokasi EWS dan Perangkap Hama", width_inch=6.5)

    add_header_styled(doc, "3.2 Memeriksa Informasi Rinci Tiang EWS", level=2)
    add_paragraph_styled(doc, "Klik salah satu nama titik (misalnya EWS-01). Layar akan menampilkan bola kemiringan lereng, grafik hujan, dan kondisi baterai aki.")
    add_image_with_caption(doc, "screenshot_ews_detail.png", "Instrumen Kemiringan Bola Inclinometer dan Riwayat Sensor", width_inch=6.5)

    add_header_styled(doc, "3.3 Memantau Hasil Perangkap Hama (Untuk Petani)", level=2)
    add_paragraph_styled(doc, "Bagi para petani, menu Perangkap Hama memberikan laporan berapa banyak hama serangga yang berhasil ditangkap setiap harinya untuk membantu rencana perawatan sawah.")
    add_image_with_caption(doc, "screenshot_traps.png", "Laporan Hasil Tangkapan Perangkap Hama Bertenaga Surya", width_inch=6.5)

    # Section 4
    add_header_styled(doc, "4. Menerima Notifikasi Darurat Telegram di HP", level=1)
    add_paragraph_styled(doc, "Sistem dilengkapi robot pengirim pesan otomatis ke grup Telegram desa:")
    add_paragraph_styled(doc, "Pastikan warga telah bergabung ke grup resmi Siaga Bencana Desa di aplikasi Telegram.", bold_prefix="• Bergabung ke Grup: ")
    add_paragraph_styled(doc, "Setiap kali tiang EWS mendeteksi hujan sangat lebat atau lereng miring, pesan bahaya akan langsung masuk ke HP seluruh warga lengkap dengan nama titik lokasi dan himbauan evakuasi.", bold_prefix="• Pesan Peringatan Otomatis: ")

    # Section 5
    add_header_styled(doc, "5. Langkah Penyelamatan Diri Saat Bahaya Longsor", level=1)
    add_paragraph_styled(doc, "1. Tetap tenang dan jangan panik. Segera bangunkan atau beritahu keluarga di dalam rumah.\n2. Jika sempat, matikan kompor gas dan matikan saklar listrik utama rumah.\n3. Bawa tas siaga bencana yang berisi surat berharga, obat pribadi, senter, dan HP.\n4. Segera keluar rumah dan ikuti jalur evakuasi menuju tanah lapang terbuka atau Balai Desa.\n5. Jangan kembali ke dalam rumah sampai petugas BPBD atau aparat desa menyatakan situasi telah aman.")

    output_path = os.path.join(DOCS_DIR, "Panduan_Penggunaan_Masyarakat_dan_Operator.docx")
    doc.save(output_path)
    print("Saved User Guide Doc to:", output_path)

if __name__ == "__main__":
    generate_technical_doc()
    generate_user_guide_doc()
